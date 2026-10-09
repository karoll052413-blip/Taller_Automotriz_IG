import os
from datetime import datetime
from tkinter import filedialog, messagebox

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle


class ReportesService:

    @staticmethod
    def _obtener_ruta_guardado(extension_defecto, nombre_sugerido):
        tipos = [("Archivo PDF", "*.pdf")] if extension_defecto == ".pdf" else [("Archivo Excel", "*.xlsx")]
        ruta = filedialog.asksaveasfilename(
            defaultextension=extension_defecto,
            filetypes=tipos,
            initialfile=f"{nombre_sugerido}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{extension_defecto}"
        )
        return ruta

    @classmethod
    def exportar_excel(cls, titulo_reporte, columnas, datos, nombre_archivo="Reporte"):
        ruta = cls._obtener_ruta_guardado(".xlsx", nombre_archivo)
        if not ruta:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte"

        fill_header = PatternFill(start_color="D32F2F", end_color="D32F2F", fill_type="solid")
        font_header = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        font_titulo = Font(name="Calibri", size=16, bold=True, color="111827")
        align_center = Alignment(horizontal="center", vertical="center")

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
        cell_t = ws.cell(row=1, column=1, value=titulo_reporte.upper())
        cell_t.font = font_titulo
        cell_t.alignment = align_center
        ws.row_dimensions[1].height = 30

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columnas))
        ws.cell(row=2, column=1, value=f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}").alignment = align_center
        ws.row_dimensions[2].height = 20

        for col_num, col_title in enumerate(columnas, 1):
            cell = ws.cell(row=4, column=col_num, value=col_title)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center

        ws.row_dimensions[4].height = 25

        for row_num, fila in enumerate(datos, 5):
            for col_num, valor in enumerate(fila, 1):
                cell = ws.cell(row=row_num, column=col_num, value=str(valor) if valor is not None else "")
                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        try:
            wb.save(ruta)
            messagebox.showinfo("Exportación Exitosa 📊", f"Archivo Excel guardado en:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error al Guardar ❌", f"No se pudo guardar el archivo Excel:\n{e}")

    @classmethod
    def exportar_pdf(cls, titulo_reporte, columnas, datos, nombre_archivo="Reporte"):
        ruta = cls._obtener_ruta_guardado(".pdf", nombre_archivo)
        if not ruta:
            return

        doc = SimpleDocTemplate(ruta, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elementos = []
        styles = getSampleStyleSheet()

        style_title = ParagraphStyle(
            'TitleStyle', parent=styles['Heading1'], fontName='Helvetica-Bold',
            fontSize=18, textColor=colors.HexColor("#D32F2F"), spaceAfter=6, alignment=1
        )
        style_sub = ParagraphStyle(
            'SubTitleStyle', parent=styles['Normal'], fontName='Helvetica',
            fontSize=10, textColor=colors.HexColor("#4B5563"), spaceAfter=15, alignment=1
        )
        style_cell = ParagraphStyle('CellStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=11)
        style_header_cell = ParagraphStyle('HeaderCellStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, textColor=colors.white, leading=12)

        elementos.append(Paragraph("TALLER AUTOMOTRIZ I.G", style_title))
        elementos.append(Paragraph(f"<b>{titulo_reporte}</b><br/>Fecha de emisión: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", style_sub))

        tabla_datos = [[Paragraph(col, style_header_cell) for col in columnas]]
        for fila in datos:
            tabla_datos.append([Paragraph(str(val) if val is not None else "", style_cell) for val in fila])

        t = Table(tabla_datos, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#D32F2F")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ]))

        elementos.append(t)

        try:
            doc.build(elementos)
            messagebox.showinfo("Exportación Exitosa 📄", f"Documento PDF generado en:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error al Guardar ❌", f"No se pudo generar el archivo PDF:\n{e}")